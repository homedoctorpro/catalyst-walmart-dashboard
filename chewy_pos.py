#!/usr/bin/env python3
"""
Chewy POS report (the $ / Units / ASP grid built with Jeff)
===========================================================
Keeps a per-SKU monthly store of Chewy retail $ (net sales) and units, merges in
whatever Chewy has sent, rebuilds the workbook through the latest complete
month, and emails it back to whoever sent the new files.

Sources, in the order they are applied:
  1. pos_store.json            history (seeded once from an existing report)
  2. Chewy sales files          .csv / .xlsx with PRODUCT_PART_NUMBER, Units Sold,
                                Net Sales and CALENDAR_MONTH (or MONTH). Only ADDS
                                months the store doesn't have $ for; a month whose
                                units disagree with what we already know (a
                                partial pull) is skipped.
  3. Brand Snapshot PDFs        units only, for months the store lacks. Poop Bags
                                are never in the top 10, so their units are the
                                brand total minus the other SKUs (flagged yellow).

Commands:
  python chewy_pos.py seed  --workbook REPORT.xlsx --data-dir _data/chewy
  python chewy_pos.py build --data-dir _data/chewy --out-dir _data/chewy/pos_reports
  python chewy_pos.py email --report X.xlsx --summary S.json --to a@b.com [--dry-run]
"""
import argparse
import csv
import datetime as dt
import glob
import json
import os
import re
import smtplib
import ssl
import sys
from collections import defaultdict
from email.message import EmailMessage

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.pagebreak import Break

import chewy_snapshots

STORE_NAME = "pos_store.json"
SKUS = [
    ("241757", "CT01", "Catalyst", "Healthy Formula Cat Litter, 10-lb bag"),
    ("241758", "CT02", "Catalyst", "Healthy Formula Cat Litter, 20-lb bag"),
    ("965502", "CT08", "Catalyst", "Healthy Formula Cat Litter, 30-lb bag"),
    ("241763", "CT05", "Catalyst", "Multi-Cat Formula Cat Litter, 10-lb bag"),
    ("241764", "CT06", "Catalyst", "Multi-Cat Formula Cat Litter, 20-lb bag"),
    ("1633142", "CT19", "Catalyst", "Pine Pellet Unscented Non-Clumping Wood Cat Litter, 20-lb bag"),
    ("1685430", "CT18", "Catalyst", "Poop Bags, 60 count"),
    ("1665670", "CT09", "Catalyst", "Sisal Cat Litter Mat, 24-in x 36-in"),
    ("241760", "CT03", "Catalyst", "Unscented Formula Cat Litter, 10-lb bag"),
    ("241761", "CT04", "Catalyst", "Unscented Formula Cat Litter, 20-lb bag"),
    ("1674830", "FW01", "Catalyst", "Wide Slatted Cat Litter Scoop"),
    ("1932182", "CT22", "Feline Fresh", "Pine Pellet Unscented Non-Clumping Cat Litter, 10-lb bag"),
    ("1932190", "CT29", "Feline Fresh", "Pine Pellet Unscented Non-Clumping Cat Litter, 20-lb bag"),
    ("1932198", "CT34", "Feline Fresh", "Pine Pellet Unscented Non-Clumping Cat Litter, 40-lb bag"),
]
PARTS = [s[0] for s in SKUS]
INFO = {s[0]: s for s in SKUS}
PART_BY_CT = {s[1]: s[0] for s in SKUS}
PART_BY_CT["CT28"] = "1932190"          # older reports labelled the FF 20-lb CT28
BAGS = "1685430"
CATALYST = [p for p, _, b, _ in SKUS if b == "Catalyst"]
FELINE = [p for p, _, b, _ in SKUS if b == "Feline Fresh"]
LITTER = ["241757", "241758", "965502", "241763", "241764", "1633142", "241760", "241761"]

# Who may trigger a reply, and who is always copied.
ALLOWED_SENDERS = {"jeff@alignsalesconsulting.com", "pross@lignetics.com"}
FIRST_NAMES = {"jeff@alignsalesconsulting.com": "Jeff", "pross@lignetics.com": "Phil"}
ALWAYS_CC = "pross@lignetics.com"
SMTP_HOST, SMTP_PORT = "smtp.gmail.com", 465


# ----------------------------------------------------------------- months
def ym_add(m, k):
    y, mo = int(m[:4]), int(m[5:7])
    t = y * 12 + mo - 1 + k
    return f"{t // 12:04d}-{t % 12 + 1:02d}"


def month_range(a, b):
    out, m = [], a
    while m <= b:
        out.append(m)
        m = ym_add(m, 1)
    return out


def fy_start(m):
    """Chewy's fiscal year starts in February."""
    y, mo = int(m[:4]), int(m[5:7])
    return f"{y if mo >= 2 else y - 1:04d}-02"


def fy_label(start):
    return f"FY{(int(start[:4])) % 100:02d}"


def mon(m):
    return dt.date(int(m[:4]), int(m[5:7]), 1).strftime("%b")


def month_long(m):
    return dt.date(int(m[:4]), int(m[5:7]), 1).strftime("%B %Y")


def to_ym(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return f"{v.year:04d}-{v.month:02d}"
    if v is None:
        return None
    mm = re.match(r"\s*(\d{4})-(\d{2})", str(v))
    return f"{mm.group(1)}-{mm.group(2)}" if mm else None


def norm_part(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


# ----------------------------------------------------------------- store
def store_path(data_dir):
    return os.path.join(data_dir, STORE_NAME)


def load_store(data_dir):
    p = store_path(data_dir)
    if not os.path.exists(p):
        raise SystemExit(f"[pos] no {STORE_NAME} in {data_dir}; run `seed` first")
    with open(p, encoding="utf-8") as f:
        s = json.load(f)
    for k in ("dollars", "units", "derived"):
        s.setdefault(k, {})
    return s


def save_store(store, data_dir):
    for part in store["derived"]:
        store["derived"][part] = sorted(set(store["derived"][part]))
    with open(store_path(data_dir), "w", encoding="utf-8") as f:
        json.dump(store, f, indent=1, sort_keys=True)


def get(store, key, part, m):
    return store[key].get(part, {}).get(m)


def is_derived(store, part, m):
    return m in store["derived"].get(part, [])


# ----------------------------------------------------------------- seed
def _is_yellow(cell):
    try:
        return str(cell.fill.fgColor.rgb).upper().endswith("FFF2CC")
    except Exception:
        return False


def seed(workbook, data_dir):
    wb = openpyxl.load_workbook(workbook, data_only=True)
    ws = wb["Summary"]
    store = {"dollars": {}, "units": {}, "derived": {},
             "seeded_from": os.path.basename(workbook)}
    blocks = []
    for r in range(1, ws.max_row + 1):
        t = ws.cell(r, 5).value
        if isinstance(t, str) and t.strip() in ("Retail Sales $", "Retail Sales Units"):
            blocks.append((t.strip(), r))
    if len(blocks) != 2:
        raise SystemExit(f"[pos] couldn't find the $ and Units blocks in {workbook}")
    for title, r0 in blocks:
        hdr = r0 + 2
        cols = {c: to_ym(ws.cell(hdr, c).value) for c in range(13, ws.max_column + 1)
                if isinstance(ws.cell(hdr, c).value, (dt.datetime, dt.date))}
        key = "dollars" if title.endswith("$") else "units"
        r = hdr + 1
        while r <= ws.max_row and str(ws.cell(r, 4).value or "").strip() != "Grand Total":
            part = PART_BY_CT.get(str(ws.cell(r, 2).value or "").strip())
            if part:
                for c, m in cols.items():
                    cell = ws.cell(r, c)
                    v = cell.value
                    if isinstance(v, (int, float)) and not isinstance(v, bool):
                        if key == "dollars":
                            store["dollars"].setdefault(part, {})[m] = round(float(v), 4)
                        else:
                            store["units"].setdefault(part, {})[m] = int(round(v))
                            if _is_yellow(cell):
                                store["derived"].setdefault(part, []).append(m)
            r += 1
    os.makedirs(data_dir, exist_ok=True)
    save_store(store, data_dir)
    n_d = sum(len(v) for v in store["dollars"].values())
    n_u = sum(len(v) for v in store["units"].values())
    print(f"[pos] seeded {STORE_NAME} from {os.path.basename(workbook)}: "
          f"{n_d} $ cells, {n_u} unit cells")
    return store


# ----------------------------------------------------------------- ingest
def read_table(rows):
    """Aggregate a Chewy sales table -> {(part, month): [dollars, units, has_dollars]}."""
    if not rows:
        return None
    hdr = [str(h).strip() if h is not None else "" for h in rows[0]]
    idx = {h: i for i, h in enumerate(hdr)}
    mcol = next((idx[c] for c in ("CALENDAR_MONTH", "MONTH") if c in idx), None)
    if mcol is None or "PRODUCT_PART_NUMBER" not in idx or "Units Sold" not in idx:
        return None
    pcol, ucol = idx["PRODUCT_PART_NUMBER"], idx["Units Sold"]
    ncol, scol = idx.get("Net Sales"), idx.get("Source")
    agg = defaultdict(lambda: [0.0, 0, False])
    for r in rows[1:]:
        if not r or len(r) <= max(mcol, pcol, ucol):
            continue
        if scol is not None and len(r) > scol and str(r[scol] or "").lower() == "derived":
            continue
        m, part = to_ym(r[mcol]), norm_part(r[pcol])
        if not m or part not in INFO:
            continue
        a = agg[(part, m)]
        u = r[ucol]
        if u not in (None, ""):
            a[1] += int(round(float(u)))
        if ncol is not None and len(r) > ncol and r[ncol] not in (None, ""):
            a[0] += float(r[ncol])
            a[2] = True
    return agg or None


def read_data_file(path):
    if path.lower().endswith(".csv"):
        with open(path, encoding="utf-8-sig", newline="") as f:
            return read_table([row for row in csv.reader(f)])
    best = None
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        agg = read_table([list(r) for r in ws.iter_rows(values_only=True)])
        if agg and (best is None or len({m for _, m in agg}) > len({m for _, m in best})):
            best = agg
    return best


def ingest_data_file(store, path, log):
    agg = read_data_file(path)
    if not agg:
        return 0
    months = sorted({m for _, m in agg})
    totals = {m: sum(a[1] for (p, mm), a in agg.items() if mm == m) for m in months}
    added = 0
    for i, m in enumerate(months):
        parts = [p for (p, mm) in agg if mm == m]
        known = [p for p in parts if get(store, "units", p, m) is not None
                 and not is_derived(store, p, m)]
        if known:
            a = sum(agg[(p, m)][1] for p in known)
            b = sum(get(store, "units", p, m) for p in known)
            if b and abs(a - b) / b > 0.03:
                log.append(f"skip {m} from {os.path.basename(path)}: {a:,} units vs "
                           f"{b:,} already known (partial pull)")
                continue
        elif i in (0, len(months) - 1) and len(months) >= 4:
            nb = [totals[x] for x in (months[1:4] if i == 0 else months[-4:-1])]
            med = sorted(nb)[len(nb) // 2]
            if med and totals[m] < 0.8 * med:
                log.append(f"skip {m} from {os.path.basename(path)}: looks partial "
                           f"({totals[m]:,} units vs ~{med:,})")
                continue
        for p in parts:
            dollars, units, has_d = agg[(p, m)]
            if has_d and get(store, "dollars", p, m) is None:
                store["dollars"].setdefault(p, {})[m] = round(dollars, 4)
                added += 1
            if get(store, "units", p, m) is None or is_derived(store, p, m):
                store["units"].setdefault(p, {})[m] = units
                if is_derived(store, p, m):
                    store["derived"][p].remove(m)
                added += 1
    return added


def brand_total(path):
    """Latest month's brand unit total from a snapshot's 'Monthly Customer Sales
    Units' chart (Chewy prints the series twice; the last bar is the report month)."""
    import fitz  # pymupdf
    full = re.sub(r"\s+", " ", " ".join(pg.get_text() for pg in fitz.open(path)))
    if "Monthly Customer Sales Units" not in full:
        return None
    seg = full.split("Monthly Customer Sales Units", 1)[1].split("Top 10 Products", 1)[0]
    seg = re.sub(r"\d{4}-\d{2}-\d{2}", " ", seg)
    vals = [int(v.replace(",", "")) for v in
            re.findall(r"(?<![\d,.%-])\d{1,3}(?:,\d{3})*(?![\d,.%])", seg)]
    if len(vals) < 2 or len(vals) % 2 or vals[:len(vals) // 2] != vals[len(vals) // 2:]:
        return None
    return vals[len(vals) // 2 - 1]


def ingest_snapshots(store, data_dir, log):
    added = 0
    snaps = chewy_snapshots.parse_folder(data_dir)
    for s in snaps:
        m = s["month"]
        for part, u in s["sku_units"].items():
            if part in INFO and get(store, "units", part, m) is None:
                store["units"].setdefault(part, {})[m] = int(u)
                added += 1
    for s in snaps:
        m = s["month"]
        if s["brand"] != "Catalyst" or get(store, "units", BAGS, m) is not None:
            continue
        others = [get(store, "units", p, m) for p in CATALYST if p != BAGS]
        if any(v is None for v in others):
            continue
        total = brand_total(os.path.join(data_dir, s["path"]))
        if total is None:
            continue
        bags = total - sum(others)
        if 0 <= bags <= 1000:
            store["units"].setdefault(BAGS, {})[m] = bags
            store["derived"].setdefault(BAGS, []).append(m)
            log.append(f"Poop Bags {m}: {bags} (brand total {total:,} minus other SKUs)")
            added += 1
    return added


# ----------------------------------------------------------------- build
def latest_units_month(store):
    count = defaultdict(int)
    for part in PARTS:
        for m in store["units"].get(part, {}):
            count[m] += 1
    full = [m for m, n in count.items() if n >= len(PARTS) - 1]
    return max(full) if full else max(count)


def latest_dollar_month(store):
    cand = set.intersection(*[set(store["dollars"].get(p, {})) for p in LITTER])
    return max(cand) if cand else None


GRAY = PatternFill("solid", fgColor="E7E6E6")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
BAR = PatternFill("solid", fgColor="BDD7EE")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
BOLD = Font(bold=True)
HDR = Font(bold=True, size=10)
TITLE = Font(bold=True, size=14)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
TOP = Border(top=Side(style="thin", color="808080"))
FMT_D = '_("$"* #,##0_);_("$"* \\(#,##0\\);_("$"* "-"??_);_(@_)'
FMT_U = "#,##0_);(#,##0)"
FMT_A = '"$"#,##0.00_);\\("$"#,##0.00\\)'
FMT_P = "0.0%"
MCOL0 = 13


def build(store, out_dir):
    lu = latest_units_month(store)
    ld = latest_dollar_month(store)
    cur_fy = fy_start(lu)
    prior_fy = ym_add(cur_fy, -12)
    months = list(reversed(month_range(prior_fy, lu)))       # newest first
    col = {m: L(MCOL0 + i) for i, m in enumerate(months)}
    lastcol = MCOL0 + len(months) - 1

    def span(a, b):  # inclusive month span -> "X{r}:Y{r}" template (newest col first)
        return f"{col[b]}{{r}}:{col[a]}{{r}}"

    fy_months = month_range(prior_fy, ym_add(cur_fy, -1))
    u_cur = month_range(cur_fy, lu)
    u_pri = [ym_add(m, -12) for m in u_cur]
    d_cur = month_range(cur_fy, ld) if ld and ld >= cur_fy else []
    d_pri = [ym_add(m, -12) for m in d_cur]
    d_label = f"({mon(cur_fy)}-{mon(ld)})" if d_cur else "(no $ yet)"
    u_label = f"({mon(cur_fy)}-{mon(lu)})"
    fy_now, fy_prev = fy_label(cur_fy), fy_label(prior_fy)

    def missing_d(part, m):
        return get(store, "units", part, m) is not None and get(store, "dollars", part, m) is None

    incomplete = {m for m in months if any(missing_d(p, m) for p in PARTS)}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"

    def put(r, c, v, font=None, fmt=None, align=None, fill=None):
        cell = ws.cell(r, c, v)
        if font:
            cell.font = font
        if fmt:
            cell.number_format = fmt
        if align:
            cell.alignment = align
        if fill:
            cell.fill = fill
        return cell

    def header(r, title, ytd_label):
        for c in range(5, lastcol + 1):
            ws.cell(r, c).fill = BAR
        put(r, 5, title, TITLE)
        if r == 1:
            put(r, 1, "Chewy", HDR)
            put(r, 2, "Lignetics", HDR)
            put(r + 1, 1, "Number", HDR)
            put(r + 1, 2, "SKU Number", HDR)
            put(r + 1, 3, "Brand", HDR)
            put(r + 1, 4, "Description", HDR)
        put(r + 1, 10, f"{fy_now} YTD vs. {fy_prev} YTD", HDR)
        put(r + 1, MCOL0, "Month Starting", HDR)
        put(r + 2, 5, fy_prev, HDR, align=CENTER)
        put(r + 2, 7, f"{fy_now} YTD\n{ytd_label}", HDR, align=CENTER)
        put(r + 2, 8, f"{fy_prev} YTD\n{ytd_label}", HDR, align=CENTER)
        put(r + 2, 10, "Chg", HDR, align=CENTER)
        put(r + 2, 11, "%", HDR, align=CENTER)
        for m in months:
            put(r + 2, MCOL0 + months.index(m), dt.date(int(m[:4]), int(m[5:]), 1),
                HDR, "mmm-yy", CENTER)
        ws.row_dimensions[r + 2].height = 30
        return r + 3

    def sku_rows(r, key):
        rows, groups = {}, []
        for brand, parts in (("Catalyst", CATALYST), ("Feline Fresh", FELINE)):
            first = r
            for p in parts:
                _, ct, b, desc = INFO[p]
                ws.cell(r, 1, int(p))
                ws.cell(r, 2, ct)
                ws.cell(r, 3, b)
                ws.cell(r, 4, desc)
                rows[p] = r
                for m in months:
                    c = MCOL0 + months.index(m)
                    if key in ("dollars", "units"):
                        v = get(store, key, p, m)
                        cell = ws.cell(r, c, v)
                        cell.number_format = FMT_D if key == "dollars" else FMT_U
                        if key == "dollars" and missing_d(p, m):
                            cell.fill = GRAY
                        if key == "units" and v is not None and is_derived(store, p, m):
                            cell.fill = YELLOW
                r += 1
            ws.cell(r, 4, f"Total {brand}")
            groups.append((brand, r, first, r - 1))
            r += 1
        r += 1
        ws.cell(r, 4, "Grand Total")
        grand = r
        for _, tr, _, _ in groups + [(None, grand, None, None)]:
            for c in range(4, lastcol + 1):
                ws.cell(tr, c).font = BOLD
                ws.cell(tr, c).border = TOP
        return r + 1, rows, groups, grand

    def ytd_formula(r, window, blank_if_missing):
        if not window:
            return None
        rng = span(window[0], window[-1]).format(r=r)
        return f'=IF(COUNTBLANK({rng})>0,"",SUM({rng}))' if blank_if_missing else f"=SUM({rng})"

    def fill_block(kind, rows, groups, grand, dref=None, uref=None):
        fmt = {"dollars": FMT_D, "units": FMT_U}.get(kind)
        cur, pri = (d_cur, d_pri) if kind == "dollars" else (u_cur, u_pri)
        if kind in ("dollars", "units"):
            for _, tr, a, b in groups:
                for m in months:
                    rng = f"{col[m]}{a}:{col[m]}{b}"
                    f = (f'=IF(COUNTBLANK({rng})>0,"",SUM({rng}))'
                         if kind == "dollars" and m in incomplete else f"=SUM({rng})")
                    put(tr, MCOL0 + months.index(m), f, fmt=fmt)
            t1, t2 = groups[0][1], groups[1][1]
            for m in months:
                c = col[m]
                f = (f'=IF(OR({c}{t1}="",{c}{t2}=""),"",{c}{t1}+{c}{t2})'
                     if kind == "dollars" and m in incomplete else f"={c}{t1}+{c}{t2}")
                put(grand, MCOL0 + months.index(m), f, fmt=fmt)
            line_rows = [(r, [p]) for p, r in rows.items()] + \
                        [(tr, CATALYST if n == "Catalyst" else FELINE) for n, tr, _, _ in groups] + \
                        [(grand, PARTS)]
            for r, parts in line_rows:
                gaps = kind == "dollars" and any(missing_d(p, m) for p in parts for m in cur)
                put(r, 5, f"=SUM({span(fy_months[0], fy_months[-1]).format(r=r)})", fmt=fmt)
                put(r, 7, ytd_formula(r, cur, gaps), fmt=fmt)
                put(r, 8, f"=SUM({span(pri[0], pri[-1]).format(r=r)})" if pri else None, fmt=fmt)
                put(r, 10, f'=IF(G{r}="","",G{r}-H{r})', fmt=fmt)
                put(r, 11, f'=IF(G{r}="","",IFERROR(G{r}/H{r}-1,"NM"))', fmt=FMT_P)
        else:
            pairs = [(rows[p], dref[0][p], uref[0][p]) for p in rows] + \
                    [(groups[i][1], dref[1][i][1], uref[1][i][1]) for i in range(2)] + \
                    [(grand, dref[2], uref[2])]
            dsum_c = span(d_cur[0], d_cur[-1]) if d_cur else None
            dsum_p = span(d_pri[0], d_pri[-1]) if d_pri else None
            for r, dr, ur in pairs:
                for m in months:
                    c = col[m]
                    put(r, MCOL0 + months.index(m),
                        f'=IF(OR({c}{ur}="",{c}{ur}=0),"NM",IF({c}{dr}="","",{c}{dr}/{c}{ur}))',
                        fmt=FMT_A)
                put(r, 5, f'=IFERROR(E{dr}/E{ur},"NM")', fmt=FMT_A)
                if dsum_c:
                    put(r, 7, f'=IF(G{dr}="","",IFERROR(G{dr}/SUM({dsum_c.format(r=ur)}),"NM"))', fmt=FMT_A)
                    put(r, 8, f'=IFERROR(H{dr}/SUM({dsum_p.format(r=ur)}),"NM")', fmt=FMT_A)
                put(r, 10, f'=IF(G{r}="","",IF(OR(ISTEXT(G{r}),ISTEXT(H{r})),"NM",G{r}-H{r}))', fmt=FMT_A)
                put(r, 11, f'=IF(G{r}="","",IF(ISTEXT(J{r}),"NM",IFERROR(J{r}/H{r},"NM")))', fmt=FMT_P)

    r = header(1, "Retail Sales $", d_label)
    r, d_rows, d_groups, d_grand = sku_rows(r, "dollars")
    fill_block("dollars", d_rows, d_groups, d_grand)
    r += 1
    r = header(r, "Retail Sales Units", u_label)
    r, u_rows, u_groups, u_grand = sku_rows(r, "units")
    fill_block("units", u_rows, u_groups, u_grand)
    asp_start = r + 1
    r = header(asp_start, "ASP", d_label)
    r, a_rows, a_groups, a_grand = sku_rows(r, "asp")
    fill_block("asp", a_rows, a_groups, a_grand,
               dref=(d_rows, d_groups, d_grand), uref=(u_rows, u_groups, u_grand))

    derived_months = sorted({m for p in store["derived"] for m in store["derived"][p] if m in months})
    gap_months = sorted(m for m in incomplete if m >= cur_fy)
    # Months with no $ at all vs. months where only some lines are missing $.
    no_dollar_months = [m for m in u_cur if not ld or m > ld]
    partial_groups, partial_since = [], None
    for label, parts in (("Feline Fresh", FELINE),
                         ("the accessories", [p for p in CATALYST if p not in LITTER])):
        gaps_here = [m for p in parts for m in d_cur if missing_d(p, m)]
        if gaps_here:
            partial_groups.append(label)
            partial_since = min(gaps_here + ([partial_since] if partial_since else []))
    notes = ["Notes",
             f"Units through {month_long(lu)} for every SKU, from Chewy's Brand Snapshots and sales files. "
             f"Retail $ (Chewy net sales) through {month_long(ld) if ld else 'n/a'}."]
    if derived_months:
        notes.append("Poop Bags (yellow) aren't in the snapshot top 10, so their units for "
                     + ", ".join(f"{mon(m)}-{m[2:4]}" for m in derived_months)
                     + " = brand total minus the other SKUs.")
    if gap_months:
        notes.append("Gray cells: Chewy hasn't sent $ for that SKU and month yet. Totals, YTD and ASP "
                     "stay blank where a line is missing instead of showing a low number.")
    notes.append(f"$ and ASP YTD compare {d_label.strip('()')}. Units YTD compare {u_label.strip('()')}. "
                 "Chewy's fiscal year starts in February.")
    notes.append(f"Built {dt.date.today():%b %d, %Y} from files sent to ligneticsdata@gmail.com.")
    r += 1
    for i, t in enumerate(notes):
        put(r + i, 1, t, BOLD if i == 0 else Font(italic=True, size=9))

    widths = {1: 9, 2: 9, 3: 11, 4: 44, 5: 13, 6: 2, 7: 13, 8: 13, 9: 2, 10: 12, 11: 9, 12: 2}
    for c, w in widths.items():
        ws.column_dimensions[L(c)].width = w
    for i in range(len(months)):
        ws.column_dimensions[L(MCOL0 + i)].width = 11
    ws.freeze_panes = "E4"
    ws.row_breaks.append(Break(id=asp_start - 1))
    ws.page_setup.paperSize = ws.PAPERSIZE_LEGAL
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.top = 1.0
    ws.oddHeader.left.text = "Chewy"
    ws.oddHeader.center.text = "Lignetics POS"

    data = wb.create_sheet("Monthly Data")
    data.append(["MONTH", "PRODUCT_PART_NUMBER", "SKU", "BRAND", "PRODUCT_NAME",
                 "Net Sales", "Units Sold", "Source"])
    for m in sorted({m for k in ("dollars", "units") for p in store[k] for m in store[k][p]}):
        for p in PARTS:
            d, u = get(store, "dollars", p, m), get(store, "units", p, m)
            if d is None and u is None:
                continue
            _, ct, b, desc = INFO[p]
            data.append([dt.date(int(m[:4]), int(m[5:]), 1), int(p), ct, b, desc, d, u,
                         "derived" if is_derived(store, p, m) else ""])
    for row in data.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "yyyy-mm-dd"

    os.makedirs(out_dir, exist_ok=True)
    name = f"Catalyst_Feline Fresh SKU L52W Report {mon(lu)}-{lu[:4]}.xlsx"
    path = os.path.join(out_dir, name)
    wb.save(path)

    def usum(ms):
        return sum(get(store, "units", p, m) or 0 for p in PARTS for m in ms)

    summary = {
        "report": path, "file_name": name,
        "units_month": lu, "dollars_month": ld,
        "fy": fy_now, "fy_prev": fy_prev,
        "units_window": u_label.strip("()"), "dollars_window": d_label.strip("()"),
        "units_ytd": usum(u_cur), "units_ytd_prior": usum(u_pri),
        "no_dollar_months": no_dollar_months,
        "partial_dollar_groups": partial_groups,
        "partial_dollar_since": partial_since,
    }
    with open(os.path.join(out_dir, "latest_summary.json"), "w", encoding="utf-8") as f:
        json.dump(summary, f, indent=1)
    return summary


# ----------------------------------------------------------------- email
def _join(items):
    items = list(items)
    if len(items) <= 2:
        return " and ".join(items)
    return ", ".join(items[:-1]) + " and " + items[-1]


def render_body(summary, names):
    lu, ld = summary["units_month"], summary["dollars_month"]
    cur, pri = summary["units_ytd"], summary["units_ytd_prior"]
    lines = [f"Hi {names},", "",
             f"Here's the Chewy POS report updated through {month_long(lu)}, "
             "rebuilt from the files sent to ligneticsdata@gmail.com.", ""]
    para = f"Units run through {month_long(lu)} for every SKU."
    if ld:
        para += f" Retail $ runs through {month_long(ld)}."
    missing = []
    gaps = summary.get("no_dollar_months") or []
    if gaps:
        missing.append(_join([month_long(m).split()[0] for m in gaps]))
    groups = summary.get("partial_dollar_groups") or []
    since = summary.get("partial_dollar_since")
    if groups and since:
        missing.append(f"for {_join(groups)} since {month_long(ym_add(since, -1)).split()[0]}")
    if missing:
        para += (" Chewy hasn't sent $ yet for " + ", or ".join(missing)
                 + ". Those cells are gray and the totals stay blank until it does.")
    lines += [para, ""]
    if pri:
        lines += [f"{summary['fy']} YTD units ({summary['units_window']}): {cur:,} vs {pri:,} "
                  f"last year ({(cur - pri) / pri * 100:+.1f}%).", ""]
    lines += ["Send new Brand Snapshots or Chewy's monthly sales file to "
              "ligneticsdata@gmail.com anytime and an updated report comes back.", "",
              "Catalyst reporting"]
    return "\n".join(lines)


def email_report(report, summary_path, to_csv, dry_run=False):
    with open(summary_path, encoding="utf-8") as f:
        summary = json.load(f)
    raw = [a.strip().lower() for a in re.split(r"[,;\s]+", to_csv or "") if a.strip()]
    to = [a for a in dict.fromkeys(raw) if a in ALLOWED_SENDERS]
    skipped = [a for a in raw if a not in ALLOWED_SENDERS]
    if skipped:
        print(f"[pos] not replying to unlisted sender(s): {', '.join(skipped)}")
    if not to:
        print("[pos] no allowed recipient; nothing sent")
        return
    cc = [ALWAYS_CC] if ALWAYS_CC not in to else []
    names = " and ".join(FIRST_NAMES.get(a, a.split("@")[0].title()) for a in to)
    sender = os.environ.get("EMAIL_USER", "ligneticsdata@gmail.com")

    msg = EmailMessage()
    msg["Subject"] = f"Chewy POS report through {month_long(summary['units_month'])}"
    msg["From"] = sender
    msg["To"] = ", ".join(to)
    if cc:
        msg["Cc"] = ", ".join(cc)
    msg.set_content(render_body(summary, names))
    with open(report, "rb") as f:
        msg.add_attachment(f.read(), maintype="application",
                           subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           filename=os.path.basename(report))
    if dry_run:
        out = os.path.splitext(report)[0] + ".eml"
        with open(out, "wb") as f:
            f.write(bytes(msg))
        print(f"[pos] dry run, wrote {out}\n---\n{render_body(summary, names)}")
        return
    pw = os.environ.get("EMAIL_APP_PASSWORD")
    if not pw:
        raise SystemExit("[pos] EMAIL_APP_PASSWORD not set")
    with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, context=ssl.create_default_context()) as srv:
        srv.login(sender, pw)
        srv.send_message(msg)
    print(f"[pos] sent '{msg['Subject']}' to {', '.join(to + cc)}")


# ----------------------------------------------------------------- cli
def gh_output(**kv):
    p = os.environ.get("GITHUB_OUTPUT")
    if p:
        with open(p, "a", encoding="utf-8") as f:
            for k, v in kv.items():
                f.write(f"{k}={v}\n")


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("seed")
    s.add_argument("--workbook", required=True)
    s.add_argument("--data-dir", required=True)
    b = sub.add_parser("build")
    b.add_argument("--data-dir", required=True)
    b.add_argument("--out-dir", required=True)
    e = sub.add_parser("email")
    e.add_argument("--report", required=True)
    e.add_argument("--summary", required=True)
    e.add_argument("--to", required=True)
    e.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    if args.cmd == "seed":
        seed(args.workbook, args.data_dir)
    elif args.cmd == "build":
        store = load_store(args.data_dir)
        log = []
        added = 0
        for path in sorted(glob.glob(os.path.join(args.data_dir, "*.csv")) +
                           glob.glob(os.path.join(args.data_dir, "*.xlsx"))):
            try:
                added += ingest_data_file(store, path, log)
            except Exception as ex:                          # noqa: BLE001
                log.append(f"could not read {os.path.basename(path)}: {ex}")
        added += ingest_snapshots(store, args.data_dir, log)
        for line in log:
            print(f"[pos] {line}")
        save_store(store, args.data_dir)
        summary = build(store, args.out_dir)
        print(f"[pos] {added} new cells; units through {summary['units_month']}, "
              f"$ through {summary['dollars_month']} -> {summary['report']}")
        gh_output(report=summary["report"],
                  summary=os.path.join(args.out_dir, "latest_summary.json"))
    else:
        email_report(args.report, args.summary, args.to, args.dry_run)


if __name__ == "__main__":
    sys.exit(main())
